"""
Apollo Microsystems Excel Format Parser
Parses financial statements from Apollo Microsystems XLSX format and converts to standard template
"""

import pandas as pd
import openpyxl
from typing import Dict, List, Optional, Any
from datetime import datetime
import re


class ApolloExcelParser:
    """Parser for Apollo Microsystems Excel format with multi-year financial data"""

    def __init__(self, excel_path: str):
        self.excel_path = excel_path
        self.wb = openpyxl.load_workbook(excel_path, data_only=True)
        self.extracted_data = {
            'company_name': '',
            'ticker_symbol': '',
            'industry': 'IT Services',
            'current_stock_price': 0.0,
            'shares_outstanding': 0.0,
            'analysis_date': datetime.now().strftime("%Y-%m-%d"),
            'periods': [],
            'income_statement': {},
            'balance_sheet': {},
            'cash_flow': {}
        }

    def extract_years_from_header(self, row: List[Any]) -> Dict[int, str]:
        """Extract fiscal years from header row"""
        year_columns = {}

        for col_idx, cell in enumerate(row):
            if cell:
                cell_str = str(cell).strip()

                # Skip common non-year headers
                if cell_str.lower() in ['particulars', 'note no', 'note no.', 'as at', 'year ended', '']:
                    continue

                # Extract year using patterns
                year = self._extract_year(cell_str)
                if year:
                    year_columns[col_idx] = year

        return year_columns

    def _extract_year(self, text: str) -> Optional[str]:
        """Extract year from text like '31 March 2023', 'FY2023', etc."""
        if not text:
            return None

        patterns = [
            r'\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+(\d{4})',
            r'FY\s*(\d{4})',
            r'\b(20\d{2})\b',
        ]

        for pattern in patterns:
            match = re.search(pattern, str(text), re.IGNORECASE)
            if match:
                year = match.group(1) if len(match.groups()) > 0 else match.group(0)
                if re.match(r'20\d{2}', year):
                    return f"FY{year}"

        return None

    def clean_number(self, value: Any) -> Optional[float]:
        """Clean and convert value to float"""
        if value is None or value == '' or value == '-':
            return None

        if isinstance(value, (int, float)):
            return float(value)

        # Remove currency symbols, commas, spaces
        cleaned = re.sub(r'[₹$,\s]', '', str(value))

        # Handle parentheses as negative
        if '(' in cleaned and ')' in cleaned:
            cleaned = '-' + cleaned.replace('(', '').replace(')', '')

        # Remove non-numeric except . and -
        cleaned = re.sub(r'[^\d.\-]', '', cleaned)

        try:
            return float(cleaned)
        except ValueError:
            return None

    def parse_sheet(self, sheet_name: str, statement_type: str) -> Dict[str, Dict[str, float]]:
        """Parse a financial statement sheet"""
        if sheet_name not in self.wb.sheetnames:
            print(f"Warning: Sheet '{sheet_name}' not found")
            return {}

        sheet = self.wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))

        if len(rows) < 3:
            return {}

        # Find year columns from first few rows
        year_columns = {}
        for row_idx in range(min(5, len(rows))):
            years = self.extract_years_from_header(rows[row_idx])
            if years:
                year_columns.update(years)

        if not year_columns:
            print(f"Warning: No year columns found in sheet '{sheet_name}'")
            return {}

        # Update periods
        for year in year_columns.values():
            if year not in self.extracted_data['periods']:
                self.extracted_data['periods'].append(year)

        # Define metric mappings
        metric_keywords = self._get_metric_keywords(statement_type)

        extracted = {}

        # Parse data rows
        for row in rows:
            if not row or len(row) < 2:
                continue

            first_cell = str(row[0]).strip().lower() if row[0] else ""
            if not first_cell:
                continue

            # Match against keywords
            for keyword, template_name in metric_keywords.items():
                if keyword in first_cell:
                    if template_name not in extracted:
                        extracted[template_name] = {}

                    # Extract values for each year
                    for col_idx, year in year_columns.items():
                        if col_idx < len(row):
                            value = self.clean_number(row[col_idx])
                            if value is not None:
                                extracted[template_name][year] = value
                    break

        return extracted

    def _get_metric_keywords(self, statement_type: str) -> Dict[str, str]:
        """Get keyword mappings for each statement type"""

        if statement_type == 'income_statement':
            return {
                'revenue from operations': 'Revenue',
                'total revenue': 'Revenue',
                'revenue': 'Revenue',
                'net sales': 'Revenue',
                'sales': 'Revenue',

                'cost of goods sold': 'Cost of Goods Sold',
                'cost of sales': 'Cost of Goods Sold',
                'cost of materials': 'Cost of Goods Sold',

                'gross profit': 'Gross Profit',

                'employee benefits': 'Operating Expenses',
                'other expenses': 'Operating Expenses',
                'operating expenses': 'Operating Expenses',
                'administrative expenses': 'Operating Expenses',

                'ebitda': 'EBITDA',

                'depreciation': 'Depreciation',
                'amortization': 'Depreciation',

                'ebit': 'EBIT',
                'operating income': 'EBIT',
                'operating profit': 'EBIT',

                'finance cost': 'Interest Expense',
                'interest expense': 'Interest Expense',
                'interest cost': 'Interest Expense',

                'tax expense': 'Tax Expense',
                'income tax': 'Tax Expense',
                'current tax': 'Tax Expense',
                'total tax': 'Tax Expense',

                'profit for the year': 'Net Profit',
                'net profit': 'Net Profit',
                'profit after tax': 'Net Profit',
                'net income': 'Net Profit',
            }

        elif statement_type == 'balance_sheet':
            return {
                'total assets': 'Total Assets',

                'current assets': 'Current Assets',
                'total current assets': 'Current Assets',

                'cash and cash equivalents': 'Cash & Equivalents',
                'cash': 'Cash & Equivalents',
                'cash and bank': 'Cash & Equivalents',

                'inventory': 'Inventory',
                'inventories': 'Inventory',
                'stock': 'Inventory',

                'trade receivables': 'Receivables',
                'receivables': 'Receivables',
                'debtors': 'Receivables',
                'accounts receivable': 'Receivables',

                'property, plant and equipment': 'Fixed Assets',
                'fixed assets': 'Fixed Assets',
                'tangible assets': 'Fixed Assets',
                'non-current assets': 'Fixed Assets',

                'total liabilities': 'Total Liabilities',

                'current liabilities': 'Current Liabilities',
                'total current liabilities': 'Current Liabilities',

                'long-term debt': 'Long-term Debt',
                'long term borrowings': 'Long-term Debt',
                'non-current debt': 'Long-term Debt',

                'short-term debt': 'Short-term Debt',
                'short term borrowings': 'Short-term Debt',
                'current borrowings': 'Short-term Debt',

                'total equity': 'Total Equity',
                "shareholders' equity": 'Total Equity',
                'equity': 'Total Equity',

                'trade payables': 'Payables',
                'payables': 'Payables',
                'creditors': 'Payables',
                'accounts payable': 'Payables',
            }

        elif statement_type == 'cash_flow':
            return {
                'cash flow from operating activities': 'Operating Cash Flow',
                'operating cash flow': 'Operating Cash Flow',
                'net cash from operations': 'Operating Cash Flow',

                'capital expenditure': 'Capital Expenditure',
                'capex': 'Capital Expenditure',
                'purchase of fixed assets': 'Capital Expenditure',

                'free cash flow': 'Free Cash Flow',

                'dividends paid': 'Dividends Paid',
                'dividend payment': 'Dividends Paid',
            }

        return {}

    def parse_all(self) -> Dict[str, Any]:
        """Parse all financial statements from the Excel file"""

        # Try common sheet name patterns
        income_sheet_patterns = ['income statement', 'profit and loss', 'p&l', 'statement of profit']
        balance_sheet_patterns = ['balance sheet', 'statement of financial position', 'assets']
        cashflow_sheet_patterns = ['cash flow', 'statement of cash flows']

        # Find sheets
        sheet_names_lower = {name.lower(): name for name in self.wb.sheetnames}

        # Parse Income Statement
        for pattern in income_sheet_patterns:
            for sheet_lower, sheet_actual in sheet_names_lower.items():
                if pattern in sheet_lower:
                    print(f"Parsing Income Statement from: {sheet_actual}")
                    self.extracted_data['income_statement'] = self.parse_sheet(sheet_actual, 'income_statement')
                    break
            if self.extracted_data['income_statement']:
                break

        # Parse Balance Sheet
        for pattern in balance_sheet_patterns:
            for sheet_lower, sheet_actual in sheet_names_lower.items():
                if pattern in sheet_lower:
                    print(f"Parsing Balance Sheet from: {sheet_actual}")
                    self.extracted_data['balance_sheet'] = self.parse_sheet(sheet_actual, 'balance_sheet')
                    break
            if self.extracted_data['balance_sheet']:
                break

        # Parse Cash Flow
        for pattern in cashflow_sheet_patterns:
            for sheet_lower, sheet_actual in sheet_names_lower.items():
                if pattern in sheet_lower:
                    print(f"Parsing Cash Flow from: {sheet_actual}")
                    self.extracted_data['cash_flow'] = self.parse_sheet(sheet_actual, 'cash_flow')
                    break
            if self.extracted_data['cash_flow']:
                break

        # Sort periods
        self.extracted_data['periods'].sort()

        # Print summary
        print(f"\nExtraction Summary:")
        print(f"Periods: {self.extracted_data['periods']}")
        print(f"Income Statement metrics: {len(self.extracted_data['income_statement'])}")
        print(f"Balance Sheet metrics: {len(self.extracted_data['balance_sheet'])}")
        print(f"Cash Flow metrics: {len(self.extracted_data['cash_flow'])}")

        return self.extracted_data


def parse_apollo_excel(excel_path: str) -> Dict[str, Any]:
    """Convenience function to parse Apollo Microsystems Excel format"""
    parser = ApolloExcelParser(excel_path)
    return parser.parse_all()
